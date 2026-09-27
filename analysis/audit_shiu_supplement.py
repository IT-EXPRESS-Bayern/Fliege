"""Audit Shiu et al. supplementary workbook without inferring cross-release IDs.

Run with the workspace Python: python analysis/audit_shiu_supplement.py
"""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data/research_sources/shiu/supplementary_tables_1-12.xlsx"
ROOT_IDS = ROOT / "data/flywire_fafb_v783/proofread_root_ids_783.npy"
SHIU_V783 = ROOT / "data/research_sources/shiu/Completeness_783.csv"
ANNOTATIONS = ROOT / "data/flywire_annotations_v2.1.0/Supplemental_file1_neuron_annotations.tsv"
OUT = ROOT / "data/research_sources/shiu/derived"
REPORT = ROOT / "analysis/shiu_supplement_audit.md"

SUGAR = "Supplemental Table 1A Sugar Fir"
WATER = "ST 6A Water Firing Rates"
MULTI = "ST 4 Interaction betw Sugar, Wa"
GROOM = "Supplemental Table 7A Grooming "
JO_COMPARE = "Supp Table 8 JO-CE and JO-F fir"
TOTAL = "Supp Table 10 Overall predictio"
BEHAVIOR = ("Supp Table 9A Behavioral Data, ", "Supp Table 9B Water silencing")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for part in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(part)
    return digest.hexdigest()


def is_id(value: object) -> bool:
    return isinstance(value, str) and len(value) == 18 and value.isdigit() and value.startswith("720575940")


def category(sheet: str, name: str) -> str | None:
    if sheet == SUGAR and name.startswith("sugar_"):
        return "sugar_grn"
    if sheet == WATER and name.startswith("water_"):
        return "water_grn"
    if sheet == MULTI and name.startswith("bitter_"):
        return "bitter_grn"
    if sheet == MULTI and name.startswith("Ir94e_"):
        return "ir94e_grn"
    if sheet == GROOM and name.startswith("JO_"):
        return "jon_" + (name[3].lower() if name[3:4] in ("C", "E", "F") else "other")
    if sheet == SUGAR and name in ("MN9_r", "MN9_l"):
        return "feeding_output"
    if sheet == GROOM and name in ("aBN1", "aDN1_l", "aDN2_l"):
        return "grooming_circuit"
    return None


def status(cat: str, name: str, root: str, roots: set[str], shiu: set[str], a: dict | None) -> str:
    if root not in roots or root not in shiu or a is None:
        return "absent_from_v783_exact_id"
    if cat.endswith("_grn"):
        good = a["flow"] == "afferent" and a["super_class"] == "sensory" and a["cell_class"] == "gustatory"
        if name.startswith("bitter_"):
            good = good and a["cell_sub_class"] == "bitter"
        return "exact_id_broad_class_match" if good else "class_requires_review"
    if cat.startswith("jon_"):
        good = a["flow"] == "afferent" and a["super_class"] == "sensory" and a["cell_class"] == "mechanosensory"
        return "exact_id_broad_class_match" if good else "class_requires_review"
    if name == "MN9_r":
        good = a["flow"] == "efferent" and a["super_class"] == "motor" and a["side"] == "right" and a["nerve"] == "PhN"
        return "exact_id_broad_class_match" if good else "class_requires_review"
    if name.startswith("aDN"):
        if a["flow"] != "efferent" or a["super_class"] != "descending":
            return "class_requires_review"
        # Workbook names carry _l while current v783 annotation lists side=right.
        return "lateralization_requires_review" if name.endswith("_l") and a["side"] != "left" else "exact_id_broad_class_match"
    return "exact_id_name_unverified"


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    roots = {str(int(x)) for x in np.load(ROOT_IDS)}
    with SHIU_V783.open(encoding="utf-8-sig", newline="") as stream:
        shiu = {row[0] for row in list(csv.reader(stream))[1:] if row and is_id(row[0])}

    selected: list[dict] = []
    sheet_inventory: list[dict] = []
    all_ids: set[str] = set()
    for sheet in wb:
        content_rows = 0
        for row_num, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if any(v is not None for v in row):
                content_rows += 1
            for value in row[:3]:
                if is_id(value):
                    all_ids.add(value)
            if len(row) < 2 or not is_id(row[0]) or not isinstance(row[1], str):
                continue
            name = row[1].strip()
            cat = category(sheet.title, name)
            if cat:
                selected.append({
                    "category": cat,
                    "paper_release": "v630",
                    "source_sheet": sheet.title,
                    "source_excel_row": row_num,
                    "source_id_cell": f"A{row_num}",
                    "source_name_cell": f"B{row_num}",
                    "source_name": name,
                    "source_root_id": row[0],
                })
        sheet_inventory.append({"sheet": sheet.title, "excel_max_row": sheet.max_row,
                                "excel_max_column": sheet.max_column, "nonempty_rows": content_rows})

    selected_ids = {entry["source_root_id"] for entry in selected}
    annotations = {}
    with ANNOTATIONS.open(encoding="utf-8", newline="") as stream:
        for row in csv.DictReader(stream, delimiter="\t"):
            if row["root_id"] in selected_ids:
                annotations[row["root_id"]] = row
    fields = ("flow", "super_class", "cell_class", "cell_sub_class", "cell_type", "side", "nerve")
    for entry in selected:
        root = entry["source_root_id"]
        a = annotations.get(root)
        entry["v783_exact_root_present"] = root in roots
        entry["shiu_v783_completeness_exact_root_present"] = root in shiu
        entry["v783_annotation"] = {field: a[field] for field in fields} if a else None
        entry["mapping_status"] = status(entry["category"], entry["source_name"], root, roots, shiu, a)

    # These are v630 reference outputs. They are qualitative regression examples for
    # v783; their exact rates must not be asserted as v783 observations.
    refs = []
    targets = [
        (SUGAR, "720575940660219265", "sugarR_150Hz", 0, "sugar_150_to_MN9_r"),
        (SUGAR, "720575940645521262", "sugarR_150Hz", 0, "sugar_150_to_MN9_l"),
        (WATER, "720575940660219265", "waterR_160Hz", 0, "water_160_to_MN9_r"),
        (WATER, "720575940645521262", "waterR_160Hz", 0, "water_160_to_MN9_l"),
        (GROOM, "720575940630907434", "JON_All_140Hz", 0, "JON_140_to_aBN1"),
        (GROOM, "720575940616185531", "JON_All_140Hz", 0, "JON_140_to_aDN1"),
        (GROOM, "720575940629806974", "JON_All_140Hz", 0, "JON_140_to_aDN2"),
        (JO_COMPARE, "720575940630907434", "JO_CE_150_Hz", 1, "JO_CE_150_to_aBN1"),
        (JO_COMPARE, "720575940630907434", "JO_F_150_Hz", 1, "JO_F_150_to_aBN1"),
        (JO_COMPARE, "720575940616185531", "JO_CE_150_Hz", 1, "JO_CE_150_to_aDN1"),
        (JO_COMPARE, "720575940616185531", "JO_F_150_Hz", 1, "JO_F_150_to_aDN1"),
        (JO_COMPARE, "720575940629806974", "JO_CE_150_Hz", 1, "JO_CE_150_to_aDN2"),
        (JO_COMPARE, "720575940629806974", "JO_F_150_Hz", 1, "JO_F_150_to_aDN2"),
    ]
    for sheet_name, root, column, id_col, label in targets:
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        col = next((i for i, value in enumerate(rows[0]) if value == column), None)
        row_num = next((i for i, row in enumerate(rows, 1) if len(row) > id_col and row[id_col] == root), None)
        if col is None or row_num is None:
            raise ValueError(f"Missing reference cell: {label}")
        refs.append({"label": label, "paper_release": "v630", "source_sheet": sheet_name,
                     "source_excel_row": row_num, "source_column_index_1based": col + 1,
                     "source_column_header": column, "source_root_id": root,
                     "mean_firing_rate_hz": rows[row_num - 1][col],
                     "v783_exact_root_present": root in roots})

    total_rows = list(wb[TOTAL].iter_rows(values_only=True))
    benchmark = {
        "paper_release": "v630",
        "source_sheet": TOTAL,
        "all_predictions": {"source_excel_row": 12, "correct": total_rows[11][2], "total": total_rows[11][3]},
        "excluding_figure_2": {"source_excel_row": 14, "correct": total_rows[13][2], "total": total_rows[13][3]},
    }
    behavioral_counts = []
    for sheet_name in BEHAVIOR:
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        for index in range(len(rows) - 2):
            header, extended, total = rows[index:index + 3]
            if not (isinstance(header[0], str) and isinstance(extended[0], str)
                    and extended[0].startswith("Extended") and total[0] == "Total"):
                continue
            for col in range(1, min(len(header), len(extended), len(total))):
                if not (isinstance(header[col], str)
                        and isinstance(extended[col], (int, float))
                        and isinstance(total[col], (int, float)) and total[col] > 0):
                    continue
                behavioral_counts.append({
                    "source_sheet": sheet_name,
                    "experiment": header[0],
                    "condition": header[col],
                    "source_header_row": index + 1,
                    "source_extended_row": index + 2,
                    "source_total_row": index + 3,
                    "source_column_index_1based": col + 1,
                    "extended": int(extended[col]),
                    "total": int(total[col]),
                    "fraction_extended": extended[col] / total[col],
                })
    manifest = {
        "source": "https://static-content.springer.com/esm/art%3A10.1038%2Fs41586-024-07763-9/MediaObjects/41586_2024_7763_MOESM2_ESM.xlsx",
        "paper": "https://doi.org/10.1038/s41586-024-07763-9",
        "workbook_sha256": sha256(WORKBOOK),
        "v783_root_array_sha256": sha256(ROOT_IDS),
        "sheet_inventory": sheet_inventory,
        "all_unique_workbook_root_ids_first_three_columns": len(all_ids),
        "all_unique_ids_exact_in_v783": len(all_ids & roots),
        "all_unique_ids_without_exact_v783_match": len(all_ids - roots),
        "candidate_counts": dict(sorted(Counter(x["category"] for x in selected).items())),
        "mapping_status_counts": dict(sorted(Counter(x["mapping_status"] for x in selected).items())),
        "candidates": selected,
        "reference_checks": refs,
        "published_prediction_summary": benchmark,
        "behavioral_counts": behavioral_counts,
    }
    OUT.mkdir(parents=True, exist_ok=True)
    json_path = OUT / "shiu_supplement_index.json"
    json_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    counts = Counter(x["category"] for x in selected)
    statuses = Counter(x["mapping_status"] for x in selected)
    lines = [
        "# Audit: Shiu et al. Ergänzungstabellen 1–12",
        "",
        f"Quelle: [Nature-Zusatzmappe]({manifest['source']}); [Artikel]({manifest['paper']}).",
        f"Lokale Datei: `{WORKBOOK.relative_to(ROOT).as_posix()}`; SHA-256 `{manifest['workbook_sha256']}`.",
        "",
        "## Umfang und Version",
        "",
        f"Die Arbeitsmappe hat {len(sheet_inventory)} Blätter. Ein Blatt (`Table legends`) ist leer. Excel-Formatierung lässt viele Blätter bis Zeile 1000 reichen; die Tabelle unten zählt tatsächliche nichtleere Zeilen.",
        f"In den ersten drei Spalten stehen {len(all_ids)} verschiedene 18-stellige FlyWire-Root-IDs. {len(all_ids & roots)} kommen exakt in `proofread_root_ids_783.npy` vor, {len(all_ids - roots)} nicht.",
        f"Die Shiu-v783-`Completeness_783.csv` enthält {len(shiu)} IDs und ist eine Teilmenge der {len(roots)} v783-Proofread-IDs (Differenz {len(roots-shiu)}).",
        "Die Zusatzmappe gehört zu den publizierten v630-Experimenten. Ihre Zahlen sind Referenzen für v630. Gleiche Root-ID in v783 ist ein Kandidat; fehlende ID wird **nicht** über Namen oder Zeilenindex ersetzt.",
        "",
        "## Für unser Modell nutzbare Kandidaten",
        "",
        "| Gruppe | Zahl | Einsatz |",
        "|---|---:|---|",
        f"| Zucker-Sensoren (`sugar_`) | {counts['sugar_grn']} | Sensorreiz Zucker → Eingang; Tabelle 1A |",
        f"| Wasser-Sensoren (`water_`) | {counts['water_grn']} | Sensorreiz Wasser → Eingang; Tabelle 6A |",
        f"| Bitter-Sensoren (`bitter_`) | {counts['bitter_grn']} | Aversiver Geschmacksreiz; Tabelle 4 |",
        f"| Ir94e-Sensoren | {counts['ir94e_grn']} | Modellspezifischer Geschmacksreiz; tatsächlicher Reiz bleibt unsicher; Tabelle 4 |",
        f"| Johnston-Organ (`JO_`) | {sum(counts[k] for k in counts if k.startswith('jon_'))} | Antennen-Mechanosensorik; Tabelle 7A |",
        f"| MN9 (links/rechts) | {counts['feeding_output']} | Auslese für Proboscis-Heben; Tabelle 1A |",
        f"| aBN1, aDN1, aDN2 | {counts['grooming_circuit']} | Antennenputz-Schaltung; Tabelle 7A |",
        "",
        f"Status der {len(selected)} ausgewählten Tabellen-IDs: " + ", ".join(f"{k}={v}" for k,v in sorted(statuses.items())) + ".",
        "Die vollständige ID-Liste mit Quellblatt, Excel-Zeile, v783-Exaktmatch und aktueller Annotation steht in `data/research_sources/shiu/derived/shiu_supplement_index.json`.",
        "",
        "## Sicherheitsrelevante Zuordnungen",
        "",
        "- `MN9_r` (`720575940660219265`, Tabelle 1A Zeile 41) kommt exakt in v783 vor; aktuelle Annotation: efferent/motor, rechts, PhN. Als **Kandidat** für Proboscis-Auslese geeignet.",
        "- `MN9_l` (`720575940645521262`, Tabelle 1A Zeile 90) fehlt in v783. Kein automatischer Ersatz durch einen ähnlich benannten Motorneurontyp.",
        "- `aBN1` (`720575940630907434`, Tabelle 7A Zeile 173) existiert exakt in v783; aktuelle Annotation ist links/central, bestätigt aber den Namen nicht unabhängig.",
        "- `aDN1_l` (`720575940616185531`) und `aDN2_l` (`720575940629806974`) existieren exakt und sind in v783 als absteigend annotiert. Aktuelles Feld `side=right` widerspricht dem `_l` im Workbook-Namen. Vor einer links/rechts-Körpersteuerung Seitenkonvention und Zellidentität manuell prüfen.",
        "- Die meisten Sensorgruppen haben einen passenden breiten v783-Zellklasseneintrag. Das bestätigt die Rolle als Sensor-Kandidat, nicht jede konkrete Rezeptor- oder Reizpräferenz.",
        "",
        "## Reproduzierbare Modellchecks",
        "",
        "Die folgenden Werte sind **v630-Mittelraten** in Hz und dienen als publizierte Kontrollmuster. Nach Umstellung auf v783 erst prüfen, ob die Richtung des Effekts und die beteiligten Zelltypen erhalten bleiben.",
        "",
        "| Check | Workbook-Zelle | v630-Mittelrate Hz |",
        "|---|---|---:|",
    ]
    for r in refs:
        lines.append(f"| `{r['label']}` | `{r['source_sheet']}` Zeile {r['source_excel_row']}, Spalte {r['source_column_index_1based']} | {r['mean_firing_rate_hz']:.2f} |")
    lines += [
        "",
        f"Tabelle 10 meldet {benchmark['all_predictions']['correct']}/{benchmark['all_predictions']['total']} korrekte empirisch geprüfte Vorhersagen (Zeile 12); ohne Abbildung 2 {benchmark['excluding_figure_2']['correct']}/{benchmark['excluding_figure_2']['total']} (Zeile 14). Diese Werte sind keine allgemeine Fliegen-Verhaltensgenauigkeit.",
        "Tabelle 11A/B–F prüft geänderte Synapsengewichte, Hemmungsstärke und Glutamat-Vorzeichen. Diese Varianten sollten später als Sensitivitätstests laufen, nicht zur Auswahl des besten Ergebnisses auf denselben Validierungsdaten dienen.",
        f"Tabelle 9A/B enthält {len(behavioral_counts)} auslesbare Verhaltenszählwerte nach Optogenetik bzw. Silencing, mit Zähler, Nenner und Excel-Zellen im JSON. Beispiel: Tabelle 9A, 50 mM Sucrose und Bitter-GRN-Aktivierung, Versuch `Gr66a > Chrimson`: Proboscis-Erweiterung 26/30 bei Licht aus (Zeilen 2–3, Spalte 6), 3/30 bei Licht an (Spalte 7). Diese Daten liefern keine direkten Muskelkräfte.",
        "Tabelle 12 listet Laborressourcen, keine zusätzlichen Modellparameter.",
        "",
        "## Blattinventar",
        "",
        "| Blatt | Nichtleere Zeilen | Excel-Maximalzeile | Spalten |",
        "|---|---:|---:|---:|",
    ]
    for s in sheet_inventory:
        lines.append(f"| `{s['sheet']}` | {s['nonempty_rows']} | {s['excel_max_row']} | {s['excel_max_column']} |")
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json_path)
    print(REPORT)


if __name__ == "__main__":
    main()
